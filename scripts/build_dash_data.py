"""Reads the finished workbook and writes the dashboard datasets to dashboard/datos/.
Usage: python3 scripts/build_dash_data.py salida/Entidades_publicas_poder_y_presupuesto_v6.xlsx"""
import openpyxl, json, re, sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent  # raíz del repositorio
sys.path.insert(0, str(BASE / 'scripts'))
from destino_labels import destino_para_mostrar
RES = BASE / 'datos' / 'resultados'
DASH = BASE / 'dashboard' / 'datos'
wb = openpyxl.load_workbook(sys.argv[1] if len(sys.argv) > 1 else BASE / 'salida' / 'Entidades_publicas_poder_y_presupuesto_v6.xlsx', data_only=True)
ws = wb['Hoja 1']
h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
def g(r, k):
    v = ws.cell(r, h[k]).value
    return v.strip() if isinstance(v, str) else v
def num(v):
    return round(float(v), 6) if isinstance(v, (int, float)) else None
sample = {e['ID'] for e in json.load(open(RES / 'verificacion.json', encoding='utf-8'))}
ent, des = [], []
for r in range(2, ws.max_row + 1):
    i = int(g(r, 'ID'))
    amb = g(r, 'Ámbito')
    ambito = 'Estatal' if amb.startswith('Estatal') else 'Autonómico'
    ccaa = g(r, 'Comunidad autónoma')
    ccaa = {'Comunidad de Madrid': 'Madrid', 'Región de Murcia': 'Murcia'}.get(ccaa, ccaa)
    nivel = g(r, 'Clasificación: veredicto')
    K, L_ = num(g(r, 'Presupuesto propio (EUR millones)')), g(r, 'Año del presupuesto')
    AL, AQ, AG, AX = (num(g(r, k)) for k in ['Crédito de personal (EUR millones)', 'Gasto de personal ejecutado (EUR millones)', 'Personal: personas (ocupadas)', 'Puestos aprobados / presupuestados'])
    BP = num(g(r, 'Gasto personal 2025 reconocido anual (EUR millones)'))
    AA = num(g(r, 'Importe contextual (EUR millones; NO sumar)'))
    nuc = num(g(r, 'Desglose: personal técnico/científico/inspector'))
    has_num = any(v is not None for v in [K, AL, AQ, AG, AX, BP, AA, nuc])
    verif = ('Contrastada por un segundo revisor (muestra)' if i in sample else
             'Fuente oficial citada' if has_num else 'Sin cifras')
    pend = []
    if nivel in ('Mantener', 'Reformar'):
        if K is None or L_ != 2026: pend.append('presupuesto 2026')
        if AL is None: pend.append('cap. 1')
        if AQ is None: pend.append('gasto ejecutado')
        if AG is None: pend.append('personas')
        if nuc is None: pend.append('desglose de personal')
    src = (g(r, 'Fuentes de la ficha') or '').split('\n')[0].strip()
    ent.append({
        'id': i, 'nombre': g(r, 'Nombre'), 'ambito': ambito, 'ccaa': ccaa,
        'periodo': (g(r, 'Periodo presidencial') or '').replace('Periodo ', ''),
        'clase': g(r, 'Clasificación: útil / no útil'), 'veredicto': nivel, 'funcion': g(r, 'Clasificación: función'),
        'duplicidad': g(r, 'Clasificación: duplicidad'), 'ue': g(r, 'Clasificación: exigida por la UE'),
        'tema': g(r, 'Tema funcional'), 'organismos_tema': num(g(r, 'Organismos del mismo tema (Estado + CCAA, verificados)')),
        'apoyo': g(r, 'Coste de apoyo (sin presupuesto propio): unidad, vocales y dietas'),
        'actividad': g(r, 'Clasificación: actividad reciente'), 'motivo': g(r, 'Clasificación: motivo'),
        'base_p100': g(r, 'Clasificación: base en el P100'), 'destino': g(r, 'Clasificación: destino de sus funciones'),
        'base_tipo': (lambda b: None if not b else 'Explícito en el P100' if b.startswith('Explícito') else 'Aplicando el criterio del P100' if b.startswith('Aplicando') else 'El P100 no trata la materia')(g(r, 'Clasificación: base en el P100')),
        'presupuesto_m': K, 'presupuesto_anio': int(L_) if isinstance(L_, (int, float)) else None,
        'presupuesto_base': g(r, 'Base del presupuesto numérico'), 'presupuesto_corte': g(r, 'Corte del presupuesto numérico'),
        'cap1_m': AL, 'gasto_personal_m': AQ, 'gasto_personal_corte': g(r, 'Corte de ejecución de personal'),
        'gasto_personal_2025_m': BP,
        'personas': AG, 'personas_fecha': str(g(r, 'Año del personal')) if g(r, 'Año del personal') is not None else None,
        'plazas': AX, 'contextual_m': AA, 'contextual_base': g(r, 'Base contextual'),
        'estado': g(r, 'Estado de cuantificación'), 'verificacion': verif,
        'pendiente': ', '.join(pend) if pend else None, 'fuente': src if src.startswith('https://') else None,
    })
    ent[-1]['destino_mostrar'] = destino_para_mostrar(ent[-1]['destino'])
    if nuc is not None or num(g(r, 'Altos cargos')) is not None:
        des.append({
            'id': i, 'nombre': g(r, 'Nombre'), 'ambito': ambito, 'veredicto': nivel,
            'nucleo': nuc, 'administracion': num(g(r, 'Desglose: personal de administración y apoyo')),
            'otro': num(g(r, 'Desglose: otro / no clasificable')), 'total': num(g(r, 'Desglose: total de referencia')),
            'libre_designacion': num(g(r, 'Puestos de libre designación')), 'altos_cargos': num(g(r, 'Altos cargos')),
            'eventuales': num(g(r, 'Personal eventual (confianza)')),
            'unidad_fecha': g(r, 'Desglose: unidad y fecha'), 'criterio': g(r, 'Desglose: criterio de clasificación'),
            'fuente': g(r, 'Desglose: fuente') if (g(r, 'Desglose: fuente') or '').startswith('https://') else None,
        })
csic = [
    {'categoria': 'Investigadores', 'personas': 5433, 'grupo': 'Científicos'},
    {'categoria': 'Investigadores en formación (predoctorales)', 'personas': 2032, 'grupo': 'Científicos'},
    {'categoria': 'Técnicos de apoyo a la investigación', 'personas': 7751, 'grupo': 'Técnicos'},
    {'categoria': 'Gestión, administración y servicios', 'personas': 1139, 'grupo': 'Administración'},
    {'categoria': 'Sin área asignada (organización central y otros)', 'personas': 1116, 'grupo': 'Sin asignar'},
]
assert sum(c['personas'] for c in csic) == 17471
ver = [{'id': e['ID'], 'campo': e['campo'], 'valor_hoja': e['valor_hoja'], 'valor_fuente': str(e.get('valor_fuente_literal')),
        'unidad': e.get('unidad'), 'resultado': e['resultado'], 'url': e.get('url'), 'nota': e.get('nota')}
       for e in json.load(open(RES / 'verificacion.json', encoding='utf-8'))]
names = {e['id']: e['nombre'] for e in ent}
for v in ver: v['nombre'] = names[v['id']]
gasto = []
for e in json.load(open(RES / 'gasto_csic_cnmc.json', encoding='utf-8')):
    corto = 'CNMC' if e['ID'] == 24 else 'CSIC'
    for c in e['capitulos']:
        cap = c['capitulo']
        if e['ID'] == 24:
            v = c.get('importe_inicial_m')
            if v is None or cap.startswith('3+4'):
                continue
            tipo = 'De paso a terceros' if cap.startswith('7') else ('De paso a terceros' if cap.startswith('4') else 'Funcionamiento propio')
            if cap.startswith('4'):
                gasto.append({'clave': 'CNMC|4 postal', 'entidad': corto, 'concepto': 'Servicio postal universal (cap. 4)', 'importe_m': 110.0, 'tipo': 'De paso a terceros', 'base': 'Crédito inicial 2026 (PGE 2023 prorrogado)', 'destino': 'Transferencia para financiar la prestación del servicio postal universal.'})
                gasto.append({'clave': 'CNMC|4 resto', 'entidad': corto, 'concepto': 'Becas, cuotas e indemnizaciones (cap. 4)', 'importe_m': round(v - 110.0, 3), 'tipo': 'Funcionamiento propio', 'base': 'Crédito inicial 2026 (PGE 2023 prorrogado)', 'destino': 'Becas, prácticas, indemnizaciones a expresidentes y exconsejeros, cuotas internacionales.'})
                continue
            gasto.append({'clave': f'CNMC|{cap}', 'entidad': corto, 'concepto': cap, 'importe_m': v, 'tipo': tipo, 'base': 'Crédito inicial 2026 (PGE 2023 prorrogado)', 'destino': c.get('destino')})
        else:
            v = c.get('importe_m')
            if v is None:
                continue
            tipo = 'Funcionamiento propio' if cap.startswith('1') else 'Sin desglosar'
            gasto.append({'clave': f'CSIC|{cap}', 'entidad': corto, 'concepto': cap, 'importe_m': v, 'tipo': tipo, 'base': 'Crédito actualizado a 31/08/2026 (IGAE)', 'destino': c.get('destino')})
assert abs(sum(g_['importe_m'] for g_ in gasto if g_['entidad'] == 'CNMC') - 3686.434) < 0.01, sum(g_['importe_m'] for g_ in gasto if g_['entidad'] == 'CNMC')
refs = json.load(open(BASE / 'datos' / 'referencias_matriz.json', encoding='utf-8'))
for r in refs:
    r['destino_mostrar'] = destino_para_mostrar(r.get('destino'))
for n, d in [('gasto_capitulos', gasto), ('entidades', ent), ('referencias_matriz', refs),
             ('personal_desglose', des), ('csic_personal', csic), ('verificacion_muestra', ver)]:
    json.dump(d, open(DASH / f'{n}.json', 'w', encoding='utf-8'), ensure_ascii=False)
    print(n, len(d), len(json.dumps(d, ensure_ascii=False)))
from collections import Counter
print(Counter(e['verificacion'] for e in ent))
print(Counter(e['ccaa'] for e in ent))
print(sum(1 for e in ent if e['pendiente']))

"""Exports a flat CSV of the classification (one row per entity) for quick reading on GitHub or in any spreadsheet.
Usage: python3 scripts/export_csv.py [salida/Entidades_publicas_poder_y_presupuesto_v6.xlsx]"""
import csv, sys
from pathlib import Path
import openpyxl

BASE = Path(__file__).resolve().parent.parent
src = sys.argv[1] if len(sys.argv) > 1 else BASE / 'salida' / 'Entidades_publicas_poder_y_presupuesto_v6.xlsx'
ws = openpyxl.load_workbook(src, data_only=True)['Hoja 1']
h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
cols = [('ID', 'ID'), ('Nombre', 'Nombre'), ('Ámbito', 'Ámbito'), ('Comunidad autónoma', 'Comunidad autónoma'),
        ('Periodo presidencial', 'Periodo presidencial'),
        ('Útil / no útil', 'Clasificación: útil / no útil'), ('Veredicto', 'Clasificación: veredicto'),
        ('Base en el P100', 'Clasificación: base en el P100'), ('Destino de sus funciones', 'Clasificación: destino de sus funciones'),
        ('Función', 'Clasificación: función'), ('Tema', 'Tema funcional'),
        ('Organismos del mismo tema', 'Organismos del mismo tema (Estado + CCAA, verificados)'),
        ('Exigida por la UE', 'Clasificación: exigida por la UE'), ('Motivo', 'Clasificación: motivo'),
        ('Presupuesto propio (M EUR)', 'Presupuesto propio (EUR millones)'), ('Año del presupuesto', 'Año del presupuesto'),
        ('Base del presupuesto', 'Base del presupuesto numérico'), ('Personas ocupadas', 'Personal: personas (ocupadas)'),
        ('Gasto de personal 2025 (M EUR)', 'Gasto personal 2025 reconocido anual (EUR millones)')]
out = BASE / 'salida' / 'clasificacion_entidades.csv'
with open(out, 'w', newline='', encoding='utf-8') as f:
    w = csv.writer(f)
    w.writerow([c[0] for c in cols])
    for r in range(2, ws.max_row + 1):
        row = []
        for _, k in cols:
            v = ws.cell(r, h[k]).value
            if k == 'ID':
                v = int(v)
            row.append('' if v is None else v)
        w.writerow(row)
print(out)
